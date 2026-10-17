"""Synthetic Sample Resume Generator for TalentScout Testing.

Creates 20 realistic synthetic resumes (PDF and DOCX) with fake names and data,
covering technical profiles across Indian locations (Bhopal, Indore, Pune, Bengaluru).
"""

from pathlib import Path
import docx

BASE_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DIR = BASE_DIR / "tests" / "fixtures" / "synthetic_resumes"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


SAMPLE_PROFILES = [
    {
        "name": "Rahul Saxena",
        "title": "Senior Java & Spring Boot Developer",
        "experience_years": 4.5,
        "location": "Bhopal, Madhya Pradesh",
        "email": "rahul.saxena.dev@fakeinbox.org",
        "phone": "+91 98260 11221",
        "education": "B.Tech in Computer Science, RGPV Bhopal",
        "skills": ["Java", "Spring Boot", "Microservices", "REST API", "PostgreSQL", "Docker", "Git"],
        "summary": "Experienced backend developer with 4.5 years of experience architecting microservices with Java, Spring Boot, and PostgreSQL. Immediate joiner located in Bhopal.",
        "fmt": "pdf",
    },
    {
        "name": "Priya Sharma",
        "title": "Java Backend Engineer",
        "experience_years": 3.5,
        "location": "Indore, Madhya Pradesh",
        "email": "priya.sharma.tech@fakeinbox.org",
        "phone": "+91 94250 88776",
        "education": "B.E. in Information Technology, SGSITS Indore",
        "skills": ["Java", "Spring Boot", "Hibernate", "MySQL", "Kafka", "AWS"],
        "summary": "Passionate Java developer with 3.5 years of experience developing high-scale REST APIs and event-driven architectures with Spring Boot and Kafka.",
        "fmt": "docx",
    },
    {
        "name": "Amitabh Mishra",
        "title": "Lead Full Stack Developer (Java + React)",
        "experience_years": 5.0,
        "location": "Bhopal, MP",
        "email": "amitabh.mishra@fakeinbox.org",
        "phone": "+91 98930 22334",
        "education": "MCA, MANIT Bhopal",
        "skills": ["Java", "Spring Boot", "React", "TypeScript", "Docker", "MongoDB", "REST API"],
        "summary": "5 years of experience building scalable enterprise web applications using Spring Boot microservices and modern React frontends.",
        "fmt": "pdf",
    },
    {
        "name": "Kavita Rao",
        "title": "Senior Spring Boot & Cloud Architect",
        "experience_years": 6.0,
        "location": "Indore, Madhya Pradesh",
        "email": "kavita.rao@fakeinbox.org",
        "phone": "+91 97550 99881",
        "education": "B.Tech, IIT Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "Kubernetes", "AWS", "Kafka", "CI/CD"],
        "summary": "Cloud native engineer with 6 years of experience optimizing distributed architectures, Spring Boot services, and containerized deployments on Kubernetes.",
        "fmt": "docx",
    },
    {
        "name": "Deepak Malviya",
        "title": "Java Developer",
        "experience_years": 2.5,
        "location": "Bhopal, MP",
        "email": "deepak.malviya@fakeinbox.org",
        "phone": "+91 98270 44556",
        "education": "BCA / MCA, LNCT Bhopal",
        "skills": ["Java", "Spring Boot", "SQL", "REST API", "Hibernate"],
        "summary": "Junior Java developer with 2.5 years of hands-on experience in building backend services, debugging SQL queries, and integrating APIs.",
        "fmt": "pdf",
    },
    {
        "name": "Sunita Jain",
        "title": "Java / J2EE Specialist",
        "experience_years": 4.0,
        "location": "Pune, Maharashtra (Open to Indore / Remote)",
        "email": "sunita.jain.pune@fakeinbox.org",
        "phone": "+91 91110 33445",
        "education": "B.E. Computer Science, Pune University",
        "skills": ["Java", "Spring Boot", "Spring MVC", "Oracle", "Docker", "Git"],
        "summary": "4 years of experience delivering robust enterprise Java services with clean design patterns and unit test coverage.",
        "fmt": "docx",
    },
    {
        "name": "Manish Chouhan",
        "title": "Backend Software Engineer",
        "experience_years": 5.5,
        "location": "Indore, Madhya Pradesh",
        "email": "manish.chouhan@fakeinbox.org",
        "phone": "+91 96300 77889",
        "education": "B.Tech, Medicaps Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "Redis", "PostgreSQL", "Docker"],
        "summary": "Senior backend developer with 5.5 years of experience focusing on caching with Redis, high-throughput Spring Boot services, and PostgreSQL tuning.",
        "fmt": "pdf",
    },
    {
        "name": "Neha Agrawal",
        "title": "Python & FastAPI Developer",
        "experience_years": 3.0,
        "location": "Bhopal, MP",
        "email": "neha.agrawal@fakeinbox.org",
        "phone": "+91 94065 66778",
        "education": "B.Tech Computer Science, Oriental Bhopal",
        "skills": ["Python", "FastAPI", "Django", "PostgreSQL", "Docker", "REST API"],
        "summary": "3 years of experience developing asynchronous microservices and data ingestion pipelines with Python and FastAPI.",
        "fmt": "docx",
    },
    {
        "name": "Sandeep Patidar",
        "title": "DevOps & Cloud Engineer",
        "experience_years": 4.0,
        "location": "Indore, MP",
        "email": "sandeep.patidar@fakeinbox.org",
        "phone": "+91 98261 55667",
        "education": "B.Tech IT, IET DAVV Indore",
        "skills": ["Docker", "Kubernetes", "AWS", "CI/CD", "Terraform", "Python", "Linux"],
        "summary": "DevOps practitioner with 4 years of experience automating cloud infrastructure, Kubernetes clusters, and GitHub Actions pipelines.",
        "fmt": "pdf",
    },
    {
        "name": "Rohan Varma",
        "title": "Frontend Engineer (React / TypeScript)",
        "experience_years": 3.0,
        "location": "Bhopal, Madhya Pradesh",
        "email": "rohan.varma@fakeinbox.org",
        "phone": "+91 94251 44332",
        "education": "B.E. Computer Technology, Bhopal",
        "skills": ["React", "TypeScript", "JavaScript", "HTML/CSS", "GraphQL"],
        "summary": "Creative frontend developer with 3 years of experience building accessible, responsive, and performant web applications in React.",
        "fmt": "docx",
    },
    {
        "name": "Ananya Kulkarni",
        "title": "Java Spring Developer",
        "experience_years": 3.8,
        "location": "Indore, Madhya Pradesh",
        "email": "ananya.kulkarni@fakeinbox.org",
        "phone": "+91 98931 99882",
        "education": "B.Tech, SVVV Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "REST API", "MySQL"],
        "summary": "Dedicated software engineer with 3.8 years in Java Spring Boot development, API design, and agile team delivery.",
        "fmt": "pdf",
    },
    {
        "name": "Vikas Tiwari",
        "title": "Senior Java Architect",
        "experience_years": 8.0,
        "location": "Bhopal, MP",
        "email": "vikas.tiwari@fakeinbox.org",
        "phone": "+91 97551 22331",
        "education": "M.Tech in Software Systems, RGPV Bhopal",
        "skills": ["Java", "Spring Boot", "Microservices", "Kafka", "AWS", "Docker", "Kubernetes"],
        "summary": "8 years of progressive experience leading backend engineering teams, decomposing monoliths into Spring Boot microservices.",
        "fmt": "docx",
    },
    {
        "name": "Shweta Rathore",
        "title": "Software Developer - Java",
        "experience_years": 4.2,
        "location": "Indore, MP",
        "email": "shweta.rathore@fakeinbox.org",
        "phone": "+91 98271 88772",
        "education": "B.E. CS, Acropolis Indore",
        "skills": ["Java", "Spring Boot", "Hibernate", "PostgreSQL", "REST API"],
        "summary": "Backend specialist with 4.2 years of experience in enterprise application development and database integration.",
        "fmt": "pdf",
    },
    {
        "name": "Gaurav Sen",
        "title": "Full Stack Engineer (Java + Angular)",
        "experience_years": 4.8,
        "location": "Bhopal, MP",
        "email": "gaurav.sen@fakeinbox.org",
        "phone": "+91 91111 66554",
        "education": "B.Tech IT, Technocrats Bhopal",
        "skills": ["Java", "Spring Boot", "Angular", "TypeScript", "MySQL", "Docker"],
        "summary": "4.8 years of full stack experience across Angular SPAs and Spring Boot microservices for fintech clients.",
        "fmt": "docx",
    },
    {
        "name": "Meera Nambiar",
        "title": "Data Scientist / Machine Learning",
        "experience_years": 3.0,
        "location": "Bengaluru (Open to Relocate)",
        "email": "meera.nambiar@fakeinbox.org",
        "phone": "+91 96301 44331",
        "education": "M.Sc Data Science, Christ University",
        "skills": ["Python", "Machine Learning", "SQL", "Pandas", "Scikit-Learn"],
        "summary": "Data scientist with 3 years experience building predictive models and data extraction algorithms in Python.",
        "fmt": "pdf",
    },
    {
        "name": "Abhishek Dubey",
        "title": "Java Spring Engineer",
        "experience_years": 5.0,
        "location": "Indore, Madhya Pradesh",
        "email": "abhishek.dubey@fakeinbox.org",
        "phone": "+91 94066 33221",
        "education": "B.Tech CS, IPS Academy Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "Docker", "REST API", "Kafka"],
        "summary": "5 years of experience developing mission-critical Java microservices with Spring Boot and Kafka messaging.",
        "fmt": "docx",
    },
    {
        "name": "Divya Chaurasia",
        "title": "Backend Developer",
        "experience_years": 3.2,
        "location": "Bhopal, MP",
        "email": "divya.chaurasia@fakeinbox.org",
        "phone": "+91 98262 11009",
        "education": "B.E. IT, SIRT Bhopal",
        "skills": ["Java", "Spring Boot", "Hibernate", "MySQL", "Git"],
        "summary": "3.2 years of experience in Java development, REST API creation, and database schema management.",
        "fmt": "pdf",
    },
    {
        "name": "Rajesh Nema",
        "title": "Senior Software Engineer",
        "experience_years": 6.5,
        "location": "Indore, MP",
        "email": "rajesh.nema@fakeinbox.org",
        "phone": "+91 94252 88990",
        "education": "B.Tech CS, SGSITS Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "AWS", "Docker", "Redis"],
        "summary": "Senior engineer with 6.5 years experience delivering cloud-ready Spring Boot applications with AWS infrastructure.",
        "fmt": "docx",
    },
    {
        "name": "Pooja Bisen",
        "title": "Java Web Developer",
        "experience_years": 4.0,
        "location": "Bhopal, MP",
        "email": "pooja.bisen@fakeinbox.org",
        "phone": "+91 98932 77665",
        "education": "MCA, Barkatullah University Bhopal",
        "skills": ["Java", "Spring Boot", "Spring MVC", "PostgreSQL", "REST API"],
        "summary": "4 years of experience writing clean, maintainable Java code and Spring Boot REST APIs.",
        "fmt": "pdf",
    },
    {
        "name": "Tarun Solanki",
        "title": "Software Engineer",
        "experience_years": 3.6,
        "location": "Indore, Madhya Pradesh",
        "email": "tarun.solanki@fakeinbox.org",
        "phone": "+91 97552 44338",
        "education": "B.E. CS, Prestige Indore",
        "skills": ["Java", "Spring Boot", "Microservices", "Docker", "MySQL"],
        "summary": "3.6 years experience in Java Spring Boot development, automated testing, and CI/CD workflows.",
        "fmt": "docx",
    },
]


def create_docx_resume(profile: dict, filepath: Path):
    """Generate a clean DOCX resume."""
    doc = docx.Document()
    doc.add_heading(profile["name"], level=0)
    doc.add_paragraph(f"{profile['title']} | {profile['location']}")
    doc.add_paragraph(f"Email: {profile['email']} | Phone: {profile['phone']}")
    doc.add_paragraph(f"Total Experience: {profile['experience_years']} Years")

    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(profile["summary"])

    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(", ".join(profile["skills"]))

    doc.add_heading("Education", level=1)
    doc.add_paragraph(profile["education"])

    doc.save(filepath)


def create_pdf_resume(profile: dict, filepath: Path):
    """Generate a clean PDF resume via reportlab (or fallback to docx if unavailable)."""
    if not HAS_REPORTLAB:
        create_docx_resume(profile, filepath.with_suffix(".docx"))
        return

    doc = SimpleDocTemplate(str(filepath), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor='#1e293b'
    )
    story.append(Paragraph(profile["name"], title_style))
    story.append(Paragraph(f"<b>{profile['title']}</b> | {profile['location']}", styles["Normal"]))
    story.append(Paragraph(f"Email: {profile['email']} | Phone: {profile['phone']}", styles["Normal"]))
    story.append(Paragraph(f"Total Experience: <b>{profile['experience_years']} Years</b>", styles["Normal"]))
    story.append(Spacer(1, 14))

    story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", styles["Heading2"]))
    story.append(Paragraph(profile["summary"], styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>TECHNICAL SKILLS</b>", styles["Heading2"]))
    story.append(Paragraph(", ".join(profile["skills"]), styles["Normal"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph("<b>EDUCATION</b>", styles["Heading2"]))
    story.append(Paragraph(profile["education"], styles["Normal"]))

    doc.build(story)


def generate_all():
    """Generate all 20 sample resumes."""
    created = []
    for i, prof in enumerate(SAMPLE_PROFILES):
        slug = prof["name"].lower().replace(" ", "_")
        ext = prof.get("fmt", "pdf")
        filename = f"{i+1:02d}_{slug}.{ext}"
        filepath = OUTPUT_DIR / filename

        if ext == "docx":
            create_docx_resume(prof, filepath)
        else:
            create_pdf_resume(prof, filepath)

        created.append(filepath)

    print(f"Successfully generated {len(created)} synthetic resumes in {OUTPUT_DIR}")
    return created


if __name__ == "__main__":
    generate_all()
